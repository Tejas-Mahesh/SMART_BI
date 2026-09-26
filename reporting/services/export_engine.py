import io

from django.core.files.base import ContentFile

from openpyxl import Workbook

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


# ============================================================
# COMMON HELPERS
# ============================================================

def _safe_value(value):
    """
    Convert values into formats suitable for PDF/Excel.
    """

    if value is None:
        return ""

    if isinstance(value, (dict, list)):
        return str(value)

    return value


def _report_rows(report):
    """
    Convert generated report data into simple rows.

    The report engine stores generated_data as JSON.
    This helper keeps export generation tolerant of the
    existing generated-data structure.
    """

    data = report.generated_data or {}

    rows = []

    if isinstance(data, dict):

        for key, value in data.items():

            if isinstance(value, dict):

                for sub_key, sub_value in value.items():

                    rows.append(
                        [
                            str(key),
                            str(sub_key),
                            _safe_value(sub_value),
                        ]
                    )

            elif isinstance(value, list):

                for index, item in enumerate(value, start=1):

                    rows.append(
                        [
                            str(key),
                            str(index),
                            _safe_value(item),
                        ]
                    )

            else:

                rows.append(
                    [
                        str(key),
                        "",
                        _safe_value(value),
                    ]
                )

    return rows


# ============================================================
# PDF EXPORT
# ============================================================

def generate_pdf_export(report):
    """
    Generate a PDF representation of a completed report.

    Returns:
        bytes
    """

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    elements = []

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    elements.append(
        Paragraph(
            report.name,
            styles["Title"],
        )
    )

    elements.append(
        Spacer(
            1,
            12,
        )
    )

    # --------------------------------------------------------
    # DESCRIPTION
    # --------------------------------------------------------

    if report.description:

        elements.append(
            Paragraph(
                report.description,
                styles["BodyText"],
            )
        )

        elements.append(
            Spacer(
                1,
                12,
            )
        )

    # --------------------------------------------------------
    # REPORT INFORMATION
    # --------------------------------------------------------

    information = [
        ["Dataset", report.dataset.name],
        [
            "Dataset Version",
            f"v{report.dataset_version.version_number}",
        ],
        ["Report Type", report.report_type],
        ["Status", report.status],
        [
            "Created",
            report.created_at.strftime(
                "%d %b %Y %H:%M"
            ),
        ],
    ]

    info_table = Table(
        information,
        colWidths=[
            140,
            340,
        ],
    )

    info_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.lightgrey,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
            ]
        )
    )

    elements.append(info_table)

    elements.append(
        Spacer(
            1,
            20,
        )
    )

    # --------------------------------------------------------
    # GENERATED DATA
    # --------------------------------------------------------

    elements.append(
        Paragraph(
            "Report Results",
            styles["Heading2"],
        )
    )

    elements.append(
        Spacer(
            1,
            8,
        )
    )

    rows = _report_rows(report)

    if rows:

        table_data = [
            [
                "Section",
                "Metric",
                "Value",
            ]
        ]

        table_data.extend(rows)

        result_table = Table(
            table_data,
            repeatRows=1,
            colWidths=[
                130,
                150,
                200,
            ],
        )

        result_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.darkgrey,
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        8,
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                ]
            )
        )

        elements.append(result_table)

    else:

        elements.append(
            Paragraph(
                "No generated report data is available.",
                styles["BodyText"],
            )
        )

    document.build(elements)

    buffer.seek(0)

    return buffer.getvalue()


# ============================================================
# EXCEL EXPORT
# ============================================================

def generate_excel_export(report):
    """
    Generate an Excel representation of a completed report.

    Returns:
        bytes
    """

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Report"

    # --------------------------------------------------------
    # REPORT INFORMATION
    # --------------------------------------------------------

    worksheet["A1"] = "Report"
    worksheet["B1"] = report.name

    worksheet["A2"] = "Dataset"
    worksheet["B2"] = report.dataset.name

    worksheet["A3"] = "Dataset Version"
    worksheet["B3"] = (
        f"v{report.dataset_version.version_number}"
    )

    worksheet["A4"] = "Report Type"
    worksheet["B4"] = report.report_type

    worksheet["A5"] = "Status"
    worksheet["B5"] = report.status

    worksheet["A7"] = "Section"
    worksheet["B7"] = "Metric"
    worksheet["C7"] = "Value"

    # --------------------------------------------------------
    # GENERATED DATA
    # --------------------------------------------------------

    rows = _report_rows(report)

    current_row = 8

    for section, metric, value in rows:

        worksheet.cell(
            row=current_row,
            column=1,
            value=section,
        )

        worksheet.cell(
            row=current_row,
            column=2,
            value=metric,
        )

        worksheet.cell(
            row=current_row,
            column=3,
            value=value,
        )

        current_row += 1

    # --------------------------------------------------------
    # COLUMN WIDTHS
    # --------------------------------------------------------

    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 30
    worksheet.column_dimensions["C"].width = 60

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    buffer = io.BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer.getvalue()