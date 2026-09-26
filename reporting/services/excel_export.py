from io import BytesIO

from django.core.files.base import ContentFile
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models import ReportExport


# ============================================================
# HELPERS
# ============================================================

def safe_value(value):
    """
    Convert nested JSON values into Excel-safe values.
    """

    if value is None:
        return ""

    if isinstance(value, (dict, list)):
        return str(value)

    return value


def format_title(value):
    """
    Convert technical JSON keys into readable labels.
    """

    return str(value).replace("_", " ").title()


def style_header(row):
    """
    Apply header formatting.
    """

    fill = PatternFill(
        fill_type="solid",
        fgColor="1F2937",
    )

    font = Font(
        bold=True,
        color="FFFFFF",
    )

    for cell in row:

        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def apply_borders(worksheet):
    """
    Apply light borders to used cells.
    """

    thin = Side(
        style="thin",
        color="D1D5DB",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    for row in worksheet.iter_rows():

        for cell in row:
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


def auto_width(worksheet):
    """
    Automatically adjust worksheet column widths.
    """

    for column_cells in worksheet.columns:

        maximum_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            value = cell.value

            if value is None:
                continue

            maximum_length = max(
                maximum_length,
                len(str(value)),
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(maximum_length + 2, 12),
            45,
        )


# ============================================================
# EXCEL EXPORT
# ============================================================

def generate_excel_export(report):
    """
    Generate an Excel export for a completed Report.

    Returns:
        ReportExport instance
    """

    export = ReportExport.objects.create(
        report=report,
        export_type="Excel",
        status="Generating",
    )

    try:

        workbook = Workbook()

        # Remove default worksheet.
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        # ====================================================
        # REPORT SUMMARY
        # ====================================================

        summary_sheet = workbook.create_sheet(
            "Report Summary"
        )

        summary_sheet["A1"] = report.name
        summary_sheet["A1"].font = Font(
            bold=True,
            size=18,
        )

        summary_sheet["A3"] = "Report Type"
        summary_sheet["B3"] = (
            report.get_report_type_display()
        )

        summary_sheet["A4"] = "Status"
        summary_sheet["B4"] = (
            report.get_status_display()
        )

        summary_sheet["A5"] = "Dataset"
        summary_sheet["B5"] = report.dataset.name

        summary_sheet["A6"] = "Dataset Version"
        summary_sheet["B6"] = (
            report.dataset_version.version_number
        )

        summary_sheet["A7"] = "Created"
        summary_sheet["B7"] = report.created_at.strftime(
            "%d %b %Y, %H:%M"
        )

        summary_sheet["A9"] = "Description"
        summary_sheet["B9"] = report.description or ""

        for cell in summary_sheet["A3:A9"]:
            cell[0].font = Font(
                bold=True
            )

        apply_borders(summary_sheet)
        auto_width(summary_sheet)

        # ====================================================
        # GENERATED DATA
        # ====================================================

        generated_data = report.generated_data or {}

        for section_name, section_data in generated_data.items():

            sheet_name = format_title(
                section_name
            )[:31]

            # Avoid duplicate worksheet names.
            original_name = sheet_name
            counter = 2

            while sheet_name in workbook.sheetnames:

                suffix = f" {counter}"

                sheet_name = (
                    original_name[:31 - len(suffix)]
                    + suffix
                )

                counter += 1

            worksheet = workbook.create_sheet(
                sheet_name
            )

            worksheet["A1"] = format_title(
                section_name
            )

            worksheet["A1"].font = Font(
                bold=True,
                size=15,
            )

            # ------------------------------------------------
            # DICTIONARY
            # ------------------------------------------------

            if isinstance(section_data, dict):

                worksheet["A3"] = "Metric"
                worksheet["B3"] = "Value"

                style_header(
                    worksheet[3]
                )

                row_number = 4

                for key, value in section_data.items():

                    worksheet.cell(
                        row=row_number,
                        column=1,
                        value=format_title(key),
                    )

                    worksheet.cell(
                        row=row_number,
                        column=2,
                        value=safe_value(value),
                    )

                    row_number += 1

            # ------------------------------------------------
            # LIST OF DICTIONARIES
            # ------------------------------------------------

            elif (
                isinstance(section_data, list)
                and section_data
                and isinstance(
                    section_data[0],
                    dict,
                )
            ):

                headers = list(
                    section_data[0].keys()
                )

                for column_number, header in enumerate(
                    headers,
                    start=1,
                ):

                    worksheet.cell(
                        row=3,
                        column=column_number,
                        value=format_title(header),
                    )

                style_header(
                    worksheet[3]
                )

                for row_number, item in enumerate(
                    section_data,
                    start=4,
                ):

                    for column_number, header in enumerate(
                        headers,
                        start=1,
                    ):

                        worksheet.cell(
                            row=row_number,
                            column=column_number,
                            value=safe_value(
                                item.get(header)
                            ),
                        )

            # ------------------------------------------------
            # LIST OF LISTS
            # ------------------------------------------------

            elif (
                isinstance(section_data, list)
                and section_data
                and isinstance(
                    section_data[0],
                    (list, tuple),
                )
            ):

                for row_number, row_data in enumerate(
                    section_data,
                    start=3,
                ):

                    for column_number, value in enumerate(
                        row_data,
                        start=1,
                    ):

                        worksheet.cell(
                            row=row_number,
                            column=column_number,
                            value=safe_value(value),
                        )

                style_header(
                    worksheet[3]
                )

            # ------------------------------------------------
            # SCALAR
            # ------------------------------------------------

            else:

                worksheet["A3"] = "Value"

                worksheet["A3"].font = Font(
                    bold=True,
                )

                worksheet["B3"] = safe_value(
                    section_data
                )

            apply_borders(
                worksheet
            )

            auto_width(
                worksheet
            )

            worksheet.freeze_panes = "A4"

        # ====================================================
        # SAVE WORKBOOK
        # ====================================================

        buffer = BytesIO()

        workbook.save(buffer)

        excel_content = buffer.getvalue()

        buffer.close()

        filename = (
            f"{report.name.replace(' ', '_')}_"
            f"{report.id}.xlsx"
        )

        export.file.save(
            filename,
            ContentFile(excel_content),
            save=False,
        )

        export.status = "Completed"
        export.completed_at = timezone.now()
        export.error_message = ""

        export.save()

        return export

    except Exception as exc:

        export.status = "Failed"
        export.error_message = str(exc)
        export.save()

        raise