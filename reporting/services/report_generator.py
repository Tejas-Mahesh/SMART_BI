from pathlib import Path
from datetime import datetime

import pandas as pd

from django.conf import settings
from django.core.files.base import ContentFile

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# SAFE HELPERS
# ============================================================

def safe_float(value, default=0.0):
    try:
        value = float(value)

        if pd.isna(value):
            return default

        return value

    except (TypeError, ValueError):
        return default


def format_number(value):
    value = safe_float(value)

    if abs(value) >= 1_000_000:
        return f"{value / 1_000_000:.2f}M"

    if abs(value) >= 1_000:
        return f"{value / 1_000:.2f}K"

    return f"{value:,.2f}"


def detect_column(df, candidates):
    """
    Find a column using flexible name matching.
    """

    normalized = {
        str(column).strip().lower().replace(" ", "_"): column
        for column in df.columns
    }

    for candidate in candidates:

        key = (
            str(candidate)
            .strip()
            .lower()
            .replace(" ", "_")
        )

        if key in normalized:
            return normalized[key]

    for column in df.columns:

        column_key = (
            str(column)
            .strip()
            .lower()
            .replace(" ", "_")
        )

        for candidate in candidates:

            candidate_key = (
                str(candidate)
                .strip()
                .lower()
                .replace(" ", "_")
            )

            if candidate_key in column_key:
                return column

    return None


# ============================================================
# LOAD DATASET
# ============================================================

def load_dataframe(dataset):
    """
    Load CSV / XLSX / XLS dataset.
    """

    file_path = dataset.file.path

    suffix = Path(file_path).suffix.lower()

    if suffix == ".csv":

        return pd.read_csv(file_path)

    if suffix == ".xlsx":

        return pd.read_excel(file_path)

    if suffix == ".xls":

        return pd.read_excel(file_path)

    raise ValueError(
        "Unsupported dataset format. "
        "Only CSV, XLSX and XLS are supported."
    )


# ============================================================
# DATA ANALYSIS
# ============================================================

def analyze_dataset(df):
    """
    Generate generic business statistics from a dataset.
    """

    if df is None or df.empty:

        return {
    "rows": int(rows),
    "columns": int(columns),
    "missing": int(missing),
    "duplicates": int(duplicates),
    "quality": float(quality),
    "revenue": float(revenue),
    "quantity": float(quantity),
    "returns": float(returns),
    "average_value": float(average_value),
    "date_column": date_column,
    "sales_column": sales_column,
    "quantity_column": quantity_column,
    "return_column": return_column,
}

    df = df.copy()

    rows = len(df)

    columns = len(df.columns)

    missing = int(df.isna().sum().sum())

    duplicates = int(df.duplicated().sum())

    valid_rows = max(rows - duplicates, 0)

    quality = (
        (valid_rows / rows) * 100
        if rows
        else 0
    )

    date_column = detect_column(
        df,
        [
            "date",
            "order_date",
            "sales_date",
            "transaction_date",
            "created_at",
        ],
    )

    sales_column = detect_column(
        df,
        [
            "sales",
            "revenue",
            "amount",
            "total_sales",
            "total_amount",
            "price",
        ],
    )

    quantity_column = detect_column(
        df,
        [
            "quantity",
            "units",
            "units_sold",
            "qty",
        ],
    )

    return_column = detect_column(
        df,
        [
            "returns",
            "return",
            "returned",
            "return_count",
        ],
    )

    revenue = 0

    quantity = 0

    returns = 0

    if sales_column:

        revenue = pd.to_numeric(
            df[sales_column],
            errors="coerce",
        ).fillna(0).sum()

    if quantity_column:

        quantity = pd.to_numeric(
            df[quantity_column],
            errors="coerce",
        ).fillna(0).sum()

    if return_column:

        returns = pd.to_numeric(
            df[return_column],
            errors="coerce",
        ).fillna(0).sum()

    average_value = (
        revenue / rows
        if rows
        else 0
    )

    return {
        "rows": rows,
        "columns": columns,
        "missing": missing,
        "duplicates": duplicates,
        "quality": quality,
        "revenue": revenue,
        "quantity": quantity,
        "returns": returns,
        "average_value": average_value,
        "date_column": date_column,
        "sales_column": sales_column,
        "quantity_column": quantity_column,
        "return_column": return_column,
    }


# ============================================================
# REPORT SUMMARY
# ============================================================

def build_report_summary(dataset, report_type):
    """
    Build the business summary used by PDF and Excel.
    """

    df = load_dataframe(dataset)

    analysis = analyze_dataset(df)

    summary = {
        "dataset_name": dataset.name,
        "dataset_type": dataset.dataset_type,
        "report_type": report_type,
        "generated_at": datetime.now(),
        "rows": analysis["rows"],
        "columns": analysis["columns"],
        "missing": analysis["missing"],
        "duplicates": analysis["duplicates"],
        "quality": analysis["quality"],
        "revenue": analysis["revenue"],
        "quantity": analysis["quantity"],
        "returns": analysis["returns"],
        "average_value": analysis["average_value"],
    }

    return df, summary


# ============================================================
# PDF GENERATION
# ============================================================

def generate_pdf(report, df, summary):
    """
    Generate a professional Smart BI PDF report.
    """

    buffer_path = (
        Path(settings.MEDIA_ROOT)
        / "reports"
        / "pdf"
    )

    buffer_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    filename = (
        f"report_{report.id}_"
        f"{datetime.now().strftime('%Y%m%d%H%M%S')}.pdf"
    )

    file_path = buffer_path / filename

    document = SimpleDocTemplate(
        str(file_path),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "SmartBITitle",
        parent=styles["Title"],
        fontSize=22,
        leading=28,
        alignment=TA_CENTER,
        spaceAfter=8,
    )

    subtitle_style = ParagraphStyle(
        "SmartBISubtitle",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        alignment=TA_CENTER,
        spaceAfter=18,
    )

    heading_style = ParagraphStyle(
        "SmartBIHeading",
        parent=styles["Heading2"],
        fontSize=14,
        leading=18,
        spaceBefore=10,
        spaceAfter=8,
    )

    normal_style = ParagraphStyle(
        "SmartBINormal",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
    )

    story = []

    story.append(
        Paragraph(
            "SMART BUSINESS INTELLIGENCE",
            title_style,
        )
    )

    story.append(
        Paragraph(
            "Business Decision Intelligence Report",
            subtitle_style,
        )
    )

    story.append(
        Paragraph(
            f"<b>Report:</b> {report.title}",
            normal_style,
        )
    )

    story.append(
        Paragraph(
            f"<b>Dataset:</b> {summary['dataset_name']}",
            normal_style,
        )
    )

    story.append(
        Paragraph(
            f"<b>Dataset Type:</b> {summary['dataset_type']}",
            normal_style,
        )
    )

    story.append(
        Paragraph(
            f"<b>Report Type:</b> {summary['report_type']}",
            normal_style,
        )
    )

    story.append(
        Paragraph(
            f"<b>Generated:</b> "
            f"{summary['generated_at'].strftime('%d %b %Y %H:%M')}",
            normal_style,
        )
    )

    story.append(Spacer(1, 12))

    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Executive Summary",
            heading_style,
        )
    )

    kpi_data = [
        ["Metric", "Value"],
        ["Rows", format_number(summary["rows"])],
        ["Columns", format_number(summary["columns"])],
        ["Missing Values", format_number(summary["missing"])],
        ["Duplicate Rows", format_number(summary["duplicates"])],
        ["Data Quality", f"{summary['quality']:.2f}%"],
        ["Revenue / Sales", format_number(summary["revenue"])],
        ["Quantity / Units", format_number(summary["quantity"])],
        ["Returns", format_number(summary["returns"])],
        ["Average Value", format_number(summary["average_value"])],
    ]

    table = Table(
        kpi_data,
        colWidths=[
            80 * mm,
            70 * mm,
        ],
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#7567f8"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "FONTNAME",
                    (0, 1),
                    (-1, -1),
                    "Helvetica",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    9,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
            ]
        )
    )

    story.append(table)

    story.append(Spacer(1, 15))

    # --------------------------------------------------------
    # DATASET STRUCTURE
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Dataset Structure",
            heading_style,
        )
    )

    columns_data = [
        ["#", "Column Name", "Data Type", "Missing"],
    ]

    for index, column in enumerate(
        df.columns,
        start=1,
    ):

        missing_count = int(
            df[column].isna().sum()
        )

        columns_data.append(
            [
                str(index),
                str(column),
                str(df[column].dtype),
                str(missing_count),
            ]
        )

    # Limit PDF table to first 40 columns
    columns_data = columns_data[:41]

    column_table = Table(
        columns_data,
        repeatRows=1,
        colWidths=[
            12 * mm,
            65 * mm,
            40 * mm,
            25 * mm,
        ],
    )

    column_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#0d2033"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.grey,
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
            ]
        )
    )

    story.append(column_table)

    story.append(Spacer(1, 15))

    # --------------------------------------------------------
    # BUSINESS INSIGHTS
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Business Insights",
            heading_style,
        )
    )

    insights = []

    if summary["quality"] >= 95:

        insights.append(
            "Dataset quality is strong with minimal invalid or duplicate records."
        )

    elif summary["quality"] >= 80:

        insights.append(
            "Dataset quality is acceptable but additional cleaning may improve analytical reliability."
        )

    else:

        insights.append(
            "Dataset quality requires attention before relying heavily on analytical results."
        )

    if summary["revenue"] > 0:

        insights.append(
            f"Detected sales/revenue value is "
            f"{format_number(summary['revenue'])}."
        )

    if summary["quantity"] > 0:

        insights.append(
            f"Detected total quantity is "
            f"{format_number(summary['quantity'])}."
        )

    if summary["returns"] > 0:

        return_rate = (
            summary["returns"]
            / summary["quantity"]
            * 100
            if summary["quantity"] > 0
            else 0
        )

        insights.append(
            f"Detected return activity is "
            f"{format_number(summary['returns'])}, "
            f"representing approximately "
            f"{return_rate:.2f}% of detected units."
        )

    for insight in insights:

        story.append(
            Paragraph(
                f"• {insight}",
                normal_style,
            )
        )

        story.append(
            Spacer(1, 4)
        )

    story.append(Spacer(1, 12))

    story.append(
        Paragraph(
            "Generated by Smart Business Decision Intelligence Management System & Forecasting.",
            subtitle_style,
        )
    )

    document.build(story)

    return file_path


# ============================================================
# EXCEL GENERATION
# ============================================================

def generate_excel(report, df, summary):
    """
    Generate a professional Excel report.
    """

    output_directory = (
        Path(settings.MEDIA_ROOT)
        / "reports"
        / "excel"
    )

    output_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    filename = (
        f"report_{report.id}_"
        f"{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    )

    file_path = output_directory / filename

    workbook = Workbook()

    # --------------------------------------------------------
    # SUMMARY SHEET
    # --------------------------------------------------------

    summary_sheet = workbook.active

    summary_sheet.title = "Executive Summary"

    summary_sheet["A1"] = (
        "SMART BUSINESS INTELLIGENCE"
    )

    summary_sheet["A1"].font = Font(
        bold=True,
        size=18,
    )

    summary_sheet["A2"] = (
        "Business Decision Intelligence Report"
    )

    summary_sheet["A2"].font = Font(
        bold=True,
        size=12,
    )

    summary_rows = [
        ("Report", report.title),
        ("Dataset", summary["dataset_name"]),
        ("Dataset Type", summary["dataset_type"]),
        ("Report Type", summary["report_type"]),
        (
            "Generated At",
            summary["generated_at"].strftime(
                "%d %b %Y %H:%M"
            ),
        ),
        ("Rows", summary["rows"]),
        ("Columns", summary["columns"]),
        ("Missing Values", summary["missing"]),
        ("Duplicate Rows", summary["duplicates"]),
        ("Data Quality", summary["quality"]),
        ("Revenue / Sales", summary["revenue"]),
        ("Quantity / Units", summary["quantity"]),
        ("Returns", summary["returns"]),
        ("Average Value", summary["average_value"]),
    ]

    start_row = 4

    for index, (label, value) in enumerate(
        summary_rows,
        start=start_row,
    ):

        summary_sheet.cell(
            row=index,
            column=1,
            value=label,
        )

        summary_sheet.cell(
            row=index,
            column=2,
            value=value,
        )

        summary_sheet.cell(
            row=index,
            column=1,
        ).font = Font(
            bold=True,
        )

    summary_sheet.column_dimensions["A"].width = 25

    summary_sheet.column_dimensions["B"].width = 35

    # --------------------------------------------------------
    # DATA SHEET
    # --------------------------------------------------------

    data_sheet = workbook.create_sheet(
        "Dataset"
    )

    for column_index, column_name in enumerate(
        df.columns,
        start=1,
    ):

        cell = data_sheet.cell(
            row=1,
            column=column_index,
            value=str(column_name),
        )

        cell.font = Font(
            bold=True,
        )

        cell.alignment = Alignment(
            horizontal="center",
        )

    for row_index, row in enumerate(
        df.itertuples(
            index=False,
            name=None,
        ),
        start=2,
    ):

        for column_index, value in enumerate(
            row,
            start=1,
        ):

            data_sheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

    # Freeze headers

    data_sheet.freeze_panes = "A2"

    # Auto width

    for column_cells in data_sheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells[:100]:

            try:

                value_length = len(
                    str(cell.value)
                )

                max_length = max(
                    max_length,
                    value_length,
                )

            except Exception:

                pass

        data_sheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 10),
            40,
        )

    # --------------------------------------------------------
    # INSIGHTS SHEET
    # --------------------------------------------------------

    insight_sheet = workbook.create_sheet(
        "Insights"
    )

    insight_sheet["A1"] = (
        "Smart BI Business Insights"
    )

    insight_sheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    insight_sheet["A3"] = "Insight"

    insight_sheet["A3"].font = Font(
        bold=True,
    )

    insight_list = []

    if summary["quality"] >= 95:

        insight_list.append(
            "Dataset quality is strong."
        )

    elif summary["quality"] >= 80:

        insight_list.append(
            "Dataset quality is acceptable."
        )

    else:

        insight_list.append(
            "Dataset quality requires improvement."
        )

    if summary["revenue"] > 0:

        insight_list.append(
            f"Revenue / sales detected: "
            f"{format_number(summary['revenue'])}"
        )

    if summary["quantity"] > 0:

        insight_list.append(
            f"Quantity detected: "
            f"{format_number(summary['quantity'])}"
        )

    if summary["returns"] > 0:

        insight_list.append(
            f"Returns detected: "
            f"{format_number(summary['returns'])}"
        )

    for index, insight in enumerate(
        insight_list,
        start=4,
    ):

        insight_sheet.cell(
            row=index,
            column=1,
            value=insight,
        )

    insight_sheet.column_dimensions[
        "A"
    ].width = 100

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    workbook.save(file_path)

    return file_path


# ============================================================
# MAIN GENERATOR
# ============================================================

def generate_report_files(report):
    """
    Main reporting pipeline.

    Generates PDF, Excel, or both.
    """

    df, summary = build_report_summary(
        report.dataset,
        report.report_type,
    )

    generated_files = {
        "pdf": None,
        "excel": None,
        "summary": summary,
    }

    if report.output_format in {
        "PDF",
        "Both",
    }:

        pdf_path = generate_pdf(
            report,
            df,
            summary,
        )

        generated_files["pdf"] = pdf_path

    if report.output_format in {
        "Excel",
        "Both",
    }:

        excel_path = generate_excel(
            report,
            df,
            summary,
        )

        generated_files["excel"] = excel_path

    return generated_files